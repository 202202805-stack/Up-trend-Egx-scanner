import os
import sys
import time
import warnings
from datetime import datetime
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import requests
import yfinance as yf

# استدعاء أداة التحميل التلقائي لبيئة Google Colab إن وجدت
try:
    from google.colab import files

    IN_COLAB = True
except ImportError:
    IN_COLAB = False

warnings.filterwarnings("ignore")


# =============================================================================
# 1. دالة سحب ودمج أسعار TRADINGVIEW اللحظية
# =============================================================================
def fetch_and_merge_tv_data(symbols_list, yf_data):
    """تسحب أسعار الجلسة الحالية المباشرة من TradingView وتدمجها مع بيانات Yahoo Finance"""
    print(
        "--- 🔗 جلب البيانات اللحظية من TradingView لدمجها مع Yahoo Finance ---"
    )
    url = "https://scanner.tradingview.com/egypt/scan"
    payload = {
        "filter": [{"left": "name", "operation": "nempty"}],
        "options": {"active_symbols_only": True},
        "columns": ["name", "description", "open", "high", "low", "close"],
        "sort": {"sortBy": "name", "sortOrder": "asc"},
        "range": [0, 300],
    }
    headers = {"User-Agent": "Mozilla/5.0"}

    tv_rows = {}
    try:
        res = requests.post(url, json=payload, headers=headers, timeout=10)
        data = res.json().get("data", [])
        today_str = datetime.now().strftime("%Y-%m-%d")

        for item in data:
            sym = item["s"].replace("EGX:", "")
            d = item["d"]
            open_p = float(d[2]) if d[2] is not None else 0.0
            high_p = float(d[3]) if d[3] is not None else open_p
            low_p = float(d[4]) if d[4] is not None else open_p
            close_p = float(d[5]) if d[5] is not None else open_p

            if close_p > 0:
                tv_rows[f"{sym}.CA"] = pd.DataFrame(
                    [{
                        "Open": round(open_p, 3),
                        "High": round(high_p, 3),
                        "Low": round(low_p, 3),
                        "Close": round(close_p, 3),
                        "Volume": 100000,  # قيمة افتراضية للسيولة اللحظية
                    }],
                    index=[pd.Timestamp(today_str)],
                )
    except Exception as e:
        print(f"⚠️ متعذر سحب TradingView live data: {e}")

    merged_dict = {}
    for ticker in yf_data.columns.levels[0]:
        df_sym = yf_data[ticker].dropna(how="all").copy()
        if ticker in tv_rows:
            tv_row = tv_rows[ticker]
            if not df_sym.empty:
                last_date_str = df_sym.index[-1].strftime("%Y-%m-%d")
                if last_date_str == today_str:
                    df_sym.iloc[-1] = tv_row.iloc[0]
                else:
                    df_sym = pd.concat([df_sym, tv_row])
            else:
                df_sym = tv_row
        merged_dict[ticker] = df_sym

    return merged_dict


# =============================================================================
# CONFIGURATION CLASS
# =============================================================================
class BullFlagConfig:

    MIN_POLE_CANDLES = 3
    MAX_POLE_CANDLES = 7
    MIN_POLE_MOVE_PCT = 12.0
    MAX_POLE_MOVE_PCT = 30.0

    MIN_FLAG_CANDLES = 5
    MAX_FLAG_CANDLES = 10

    MIN_FLAG_RETRACE_PCT = 10.0
    MAX_FLAG_RETRACE_PCT = 38.2

    MIN_VOLUME_MULTIPLIER = 1.2
    FLAG_VOLUME_RATIO = 0.8
    MIN_QUALITY_SCORE = 50


# =============================================================================
# OPTIMIZED DETECTOR ENGINE
# =============================================================================
class FastBullFlagDetector:

    def __init__(self, config=None):
        self.config = config or BullFlagConfig()

    def process_stock(self, df, symbol):
        if df is None or len(df) < 50:
            return []

        df = df.copy()
        df.columns = [c.lower().replace(" ", "_") for c in df.columns]
        df = df.dropna()

        if "volume" not in df.columns or df["volume"].sum() == 0:
            return []

        cfg = self.config
        closes = df["close"].values
        highs = df["high"].values
        lows = df["low"].values
        opens = df["open"].values
        volumes = df["volume"].values
        dates = df.index.strftime("%Y-%m-%d").values
        n = len(df)
        current_price = closes[-1]

        poles = []

        for i in range(cfg.MAX_POLE_CANDLES, n - cfg.MIN_FLAG_CANDLES):
            for pole_len in range(cfg.MIN_POLE_CANDLES, cfg.MAX_POLE_CANDLES + 1):
                start_idx = i - pole_len
                end_idx = i

                if start_idx < 0:
                    continue

                p_low = lows[start_idx:end_idx].min()
                p_high = highs[start_idx:end_idx].max()
                p_open = opens[start_idx]

                if p_open <= 0:
                    continue

                pole_move_pct = ((p_high - p_open) / p_open) * 100.0

                if not (
                    cfg.MIN_POLE_MOVE_PCT <= pole_move_pct <= cfg.MAX_POLE_MOVE_PCT
                ):
                    continue

                p_closes = closes[start_idx:end_idx]
                p_opens = opens[start_idx:end_idx]
                green_ratio = np.sum(p_closes > p_opens) / pole_len

                if green_ratio < 0.5:
                    continue

                vol_start = max(0, start_idx - 20)
                avg_vol_before = (
                    np.mean(volumes[vol_start:start_idx]) if start_idx > 0 else 0
                )
                pole_avg_vol = np.mean(volumes[start_idx:end_idx])

                vol_mult = (
                    (pole_avg_vol / avg_vol_before) if avg_vol_before > 0 else 0
                )
                if vol_mult < cfg.MIN_VOLUME_MULTIPLIER:
                    continue

                x = np.arange(pole_len)
                p_highs = highs[start_idx:end_idx]
                slope = (
                    np.cov(x, p_highs)[0, 1] / np.var(x) if np.var(x) > 0 else 0
                )

                if slope <= 0:
                    continue

                poles.append({
                    "start_idx": start_idx,
                    "end_idx": end_idx - 1,
                    "pole_low": p_low,
                    "pole_high": p_high,
                    "pole_move_pct": pole_move_pct,
                    "green_ratio": green_ratio,
                    "volume_multiplier": vol_mult,
                    "pole_avg_vol": pole_avg_vol,
                })

        poles = self._remove_overlapping_poles(poles)
        if not poles:
            return []

        patterns = []
        for pole in poles:
            p_end = pole["end_idx"]
            p_range = pole["pole_high"] - pole["pole_low"]

            for flag_len in range(cfg.MIN_FLAG_CANDLES, cfg.MAX_FLAG_CANDLES + 1):
                f_start = p_end + 1
                f_end = f_start + flag_len

                if f_end >= n:
                    continue

                f_highs = highs[f_start:f_end]
                f_lows = lows[f_start:f_end]
                f_closes = closes[f_start:f_end]
                f_vols = volumes[f_start:f_end]

                f_high = np.max(f_highs)
                f_low = np.min(f_lows)

                retrace = pole["pole_high"] - f_low
                retrace_pct = (
                    (retrace / p_range) * 100.0 if p_range > 0 else 100.0
                )

                if not (
                    cfg.MIN_FLAG_RETRACE_PCT <= retrace_pct <= cfg.MAX_FLAG_RETRACE_PCT
                ) or f_high > pole["pole_high"] * 1.01:
                    continue

                pole_mid = (pole["pole_high"] + pole["pole_low"]) / 2.0
                if f_low < pole_mid:
                    continue

                x = np.arange(flag_len)
                slope = (
                    np.cov(x, f_closes)[0, 1] / np.var(x) if np.var(x) > 0 else 0
                )
                if slope > 0.001:
                    continue

                f_avg_vol = np.mean(f_vols)
                flag_vol_ratio = (
                    f_avg_vol / pole["pole_avg_vol"]
                    if pole["pole_avg_vol"] > 0
                    else 1.0
                )

                if flag_vol_ratio > cfg.FLAG_VOLUME_RATIO:
                    continue

                breakout_idx = f_end
                if breakout_idx >= n:
                    continue

                breakout_close = closes[breakout_idx]
                breakout_open = opens[breakout_idx]

                if breakout_close > f_high:
                    entry_price = f_high
                    stop_loss = f_low
                    target_price = entry_price + (p_range * 0.5)

                    if entry_price <= stop_loss:
                        continue

                    trade_res = self._simulate_trade(
                        closes,
                        highs,
                        lows,
                        dates,
                        breakout_idx + 1,
                        entry_price,
                        stop_loss,
                        target_price,
                    )

                    score = self._score_pattern(
                        pole["pole_move_pct"],
                        f_high,
                        f_low,
                        entry_price,
                        flag_vol_ratio,
                        retrace_pct,
                        breakout_close,
                        breakout_open,
                    )

                    if score >= cfg.MIN_QUALITY_SCORE:
                        patterns.append({
                            "Stock Name": symbol,
                            "Entry Date": trade_res["entry_date"],
                            "Entry Price": round(entry_price, 3),
                            "Target": round(target_price, 3),
                            "Stop Loss": round(stop_loss, 3),
                            "Exit Date": trade_res["exit_date"],
                            "Current Price": round(current_price, 3),
                            "Status": trade_res["status"],
                            "pattern_date": dates[breakout_idx],
                            "quality_score": score,
                        })

        return self._remove_overlapping_patterns(patterns)

    def _simulate_trade(
        self,
        closes,
        highs,
        lows,
        dates,
        start_idx,
        entry_price,
        stop_loss,
        target_price,
    ):
        n = len(closes)

        if start_idx >= n:
            return {
                "status": "OPEN",
                "entry_date": dates[-1],
                "exit_date": "",  # يُترك فارغاً للصفقات المفتوحة
            }

        entry_date = dates[start_idx]

        for i in range(start_idx, n):
            if lows[i] <= stop_loss:
                return {
                    "status": "LOSS",
                    "entry_date": entry_date,
                    "exit_date": dates[i],
                }

            if highs[i] >= target_price:
                return {
                    "status": "WIN",
                    "entry_date": entry_date,
                    "exit_date": dates[i],
                }

        return {
            "status": "OPEN",
            "entry_date": entry_date,
            "exit_date": "",  # يُترك فارغاً للصفقات المفتوحة
        }

    def _score_pattern(
        self,
        pole_move,
        f_high,
        f_low,
        entry_price,
        vol_ratio,
        retrace,
        b_close,
        b_open,
    ):
        score = 0
        score += 25 if pole_move >= 20 else (20 if pole_move >= 15 else 15)

        flag_range_pct = ((f_high - f_low) / entry_price) * 100.0
        score += 25 if flag_range_pct <= 2 else (15 if flag_range_pct <= 4 else 5)

        score += 20 if vol_ratio <= 0.5 else 10
        score += 15 if retrace <= 25.0 else 8

        if b_close > b_open:
            score += 15

        return min(score, 100)

    def _remove_overlapping_poles(self, items):
        if not items:
            return []
        sorted_items = sorted(
            items, key=lambda x: x["pole_move_pct"], reverse=True
        )
        result = []
        for item in sorted_items:
            overlap = any(
                not (
                    item["end_idx"] < r["start_idx"]
                    or item["start_idx"] > r["end_idx"]
                )
                for r in result
            )
            if not overlap:
                result.append(item)
        return result

    def _remove_overlapping_patterns(self, patterns):
        if not patterns:
            return []
        sorted_patterns = sorted(
            patterns, key=lambda x: x["quality_score"], reverse=True
        )
        result = []
        for p in sorted_patterns:
            overlap = any(p["pattern_date"] == r["pattern_date"] for r in result)
            if not overlap:
                result.append(p)
        return result


# =============================================================================
# FORMATTED EXCEL EXPORTER PIPELINE
# =============================================================================


def export_exact_matching_excel(df, filename):
    # الأعمدة الـ 8 المطلوبة بالكامل وبالترتيب المكون بالإنجليزي
    columns_order = [
        "Stock Name",
        "Entry Date",
        "Entry Price",
        "Target",
        "Stop Loss",
        "Exit Date",
        "Current Price",
        "Status",
    ]

    export_df = df[columns_order].copy()
    export_df = export_df.drop_duplicates(
        subset=["Stock Name", "Entry Date"], keep="first"
    )

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        export_df.to_excel(writer, sheet_name="Backtest Results", index=False)

    wb = openpyxl.load_workbook(filename)
    ws = wb["Backtest Results"]

    # إظهار خطوط الشبكة
    ws.views.sheetView[0].showGridLines = True

    # التنسيقات والألوان
    header_fill = PatternFill(
        start_color="1F497D", end_color="1F497D", fill_type="solid"
    )  # أزرق داكن
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    win_fill = PatternFill(
        start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
    )  # أخضر فاتح
    win_font = Font(name="Segoe UI", size=10, bold=True, color="375623")

    loss_fill = PatternFill(
        start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
    )  # برتقالي/أحمر فاتح
    loss_font = Font(name="Segoe UI", size=10, bold=True, color="C65911")

    open_fill = PatternFill(
        start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
    )  # أصفر فاتح
    open_font = Font(name="Segoe UI", size=10, bold=True, color="7F6000")

    data_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # تنسيق الهيدر
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # تنسيق الصفوف
    for row in range(2, ws.max_row + 1):
        status_val = df.iloc[row - 2]["Status"] if "Status" in df.columns else ""

        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

            # تنسيق الأرقام بمرتبتين عشريتين للأسعار
            if col in [3, 4, 5, 7]:
                cell.number_format = "#,##0.00"

            # تلوين الصفوف بناءً على حالة الصفقة
            if status_val == "WIN":
                cell.fill = win_fill
            elif status_val == "LOSS":
                cell.fill = loss_fill
            elif status_val == "OPEN":
                cell.fill = open_fill

    # ضبط عرض الأعمدة تلقائياً
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 15)

    wb.save(filename)


# =============================================================================
# MAIN EXECUTION PIPELINE
# =============================================================================

if __name__ == "__main__":
    egyptian_stocks = [
        "AALR.CA", "ABUK.CA", "ACAMD.CA", "ACAP.CA", "ACGC.CA", "ACTF.CA",
        "ADCI.CA", "ADIB.CA", "ADPC.CA", "ADRI.CA", "AFDI.CA", "AFMC.CA",
        "AIDC.CA", "AIFI.CA", "AIH.CA", "AJWA.CA", "ALCN.CA", "ALEX.CA",
        "ALUM.CA", "AMER.CA", "AMES.CA", "AMIA.CA", "AMII.CA", "AMOC.CA",
        "AMPI.CA", "APSW.CA", "ARAB.CA", "ARCC.CA", "AREH.CA", "ASCM.CA",
        "ASPI.CA", "ATLC.CA", "ATQA.CA", "AXPH.CA", "BIDI.CA", "BIGP.CA",
        "BINV.CA", "BIOC.CA", "BONY.CA", "BTFH.CA", "CAED.CA", "CANA.CA",
        "CCAP.CA", "CCRS.CA", "CEFM.CA", "CERA.CA", "CFGH.CA", "CICH.CA",
        "CIEB.CA", "CIRA.CA", "CLHO.CA", "CNFN.CA", "COMI.CA", "COPR.CA",
        "COSG.CA", "CPCI.CA", "CPME.CA", "CRST.CA", "CSAG.CA", "DAPH.CA",
        "DCRC.CA", "DEIN.CA", "DGTZ.CA", "DOMT.CA", "DSCW.CA", "DTPP.CA",
        "EALR.CA", "EASB.CA", "EAST.CA", "EBSC.CA", "ECAP.CA", "EDFM.CA",
        "EEII.CA", "EFIC.CA", "EFID.CA", "EFIH.CA", "EGAL.CA", "EGAS.CA",
        "EGBE.CA", "EGCH.CA", "EGREF.CA", "EGSA.CA", "EGTS.CA", "EHDR.CA",
        "ELAB.CA", "ELEC.CA", "ELKA.CA", "ELNA.CA", "ELSH.CA", "ELWA.CA",
        "EMFD.CA", "ENGC.CA", "EOSB.CA", "EPCO.CA", "EPPK.CA", "ETEL.CA",
        "ETRS.CA", "EXPA.CA", "FAIT.CA", "FAITA.CA", "FCMD.CA", "FIRE.CA",
        "FNAR.CA", "FTNS.CA", "FWRY.CA", "GBCO.CA", "GDWA.CA", "GGCC.CA",
        "GGRN.CA", "GIHD.CA", "GMCI.CA", "GOUR.CA", "GPIM.CA", "GRCA.CA",
        "GSSC.CA", "GTEX.CA", "GTHE.CA", "GTWL.CA", "HBCO.CA", "HDBK.CA",
        "HELI.CA", "HRHO.CA", "IBCT.CA", "ICFC.CA", "ICID.CA", "IDRE.CA",
        "IEEC.CA", "IFAP.CA", "INEG.CA", "INFI.CA", "IRON.CA", "ISMA.CA",
        "ISMQ.CA", "ISPH.CA", "JUFO.CA", "KABO.CA", "KORA.CA", "KRDI.CA",
        "KWIN.CA", "KZPC.CA", "LCSW.CA", "LKGP.CA", "LUTS.CA", "MAAL.CA",
        "MASR.CA", "MBEG.CA", "MBSC.CA", "MCQE.CA", "MCRO.CA", "MENA.CA",
        "MEPA.CA", "MFPC.CA", "MFSC.CA", "MHOT.CA", "MICH.CA", "MILS.CA",
        "MIPH.CA", "MOED.CA", "MOIL.CA", "MOIN.CA", "MOSC.CA", "MPCI.CA",
        "MPCO.CA", "MPRC.CA", "MTIE.CA", "NAHO.CA", "NARE.CA", "NCCW.CA",
        "NCGC.CA", "NEDA.CA", "NHPS.CA", "NINH.CA", "NIPH.CA", "OBRI.CA",
        "OCAP.CA", "OCDI.CA", "OCPH.CA", "ODIN.CA", "OFH.CA", "OIH.CA",
        "OLFI.CA", "ORAS.CA", "ORHD.CA", "ORWE.CA", "PHAR.CA", "PHDC.CA",
        "PHGC.CA", "PHTV.CA", "POUL.CA", "PRCL.CA", "PRDC.CA", "PRMH.CA",
        "QNBE.CA", "RACC.CA", "RAKT.CA", "RAYA.CA", "RKAZ.CA", "RMDA.CA",
        "RMTV.CA", "ROTO.CA", "RREI.CA", "RTVC.CA", "RUBX.CA", "SAUD.CA",
        "SCEM.CA", "SCFM.CA", "SCTS.CA", "SDTI.CA", "SEIG.CA", "SIEG.CA",
        "SIPC.CA", "SKPC.CA", "SMFR.CA", "SNFC.CA", "SPIN.CA", "SPMD.CA",
        "SUCE.CA", "SUGR.CA", "SVCE.CA", "SWDY.CA", "TALM.CA", "TANM.CA",
        "TAQA.CA", "TMGH.CA", "TORA.CA", "TWSA.CA", "TYCN.CA", "UBEE.CA",
        "UEFM.CA", "UEGC.CA", "UNIP.CA", "UNIT.CA", "UPMS.CA", "UTOP.CA",
        "VALU.CA", "VERT.CA", "VLMR.CA", "VLMRA.CA", "WCDF.CA", "WKOL.CA",
        "ZEOT.CA", "ZMID.CA"
    ]

    print("=" * 70)
    print("STARTING FAST BULK BACKTEST FOR EGYPTIAN STOCKS (1-YEAR PERIOD + LIVE TV DATA)")
    print("=" * 70)

    start_time = time.time()

    # 1. سحب بيانات سنة واحدة من Yahoo Finance
    batch_data = yf.download(
        egyptian_stocks,
        period="1y",
        interval="1d",
        group_by="ticker",
        threads=True,
        progress=False,
    )

    # 2. سحب ودمج الأسعار اللحظية من TradingView
    cleaned_symbols = [s.replace(".CA", "") for s in egyptian_stocks]
    merged_data = fetch_and_merge_tv_data(cleaned_symbols, batch_data)

    detector = FastBullFlagDetector()
    all_trades = []

    for symbol in egyptian_stocks:
        try:
            if symbol in merged_data and not merged_data[symbol].empty:
                df = merged_data[symbol].dropna(how="all")
                trades = detector.process_stock(df, symbol)
                if trades:
                    all_trades.extend(trades)
        except Exception:
            continue

    execution_time = time.time() - start_time

    if all_trades:
        results_df = pd.DataFrame(all_trades)

        # إزالة التكرارات
        results_df = results_df.drop_duplicates(
            subset=["Stock Name", "Entry Date"], keep="first"
        )

        # الترتيب الزمني من الأقدم إلى الأحدث
        results_df["Entry Date_dt"] = pd.to_datetime(results_df["Entry Date"])
        results_df = results_df.sort_values(
            by="Entry Date_dt", ascending=True
        ).reset_index(drop=True)

        excel_filename = "bull_flag_1y_live_report.xlsx"
        export_exact_matching_excel(results_df, excel_filename)

        print("\n" + "=" * 70)
        print("SUMMARY RESULTS")
        print("=" * 70)
        print(f"Total Time Taken      : {round(execution_time, 2)} seconds")
        print(f"Total Trades Found    : {len(results_df)}")
        print("=" * 70)
        print(f"Report exported to    : {excel_filename}")

        # التنزيل التلقائي في Colab
        if IN_COLAB:
            files.download(excel_filename)
    else:
        print("\nNo trades detected across the specified stocks.")
