import os
import sys
import warnings
import datetime
import requests
import openpyxl
import numpy as np
import pandas as pd
import yfinance as yf
import concurrent.futures
from dataclasses import dataclass
from typing import List, Optional, Tuple, Dict
from enum import Enum
from scipy.signal import argrelextrema
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

warnings.filterwarnings('ignore')

# ==============================================================================
# 1. قائمة الأسهم المصرية المحددة
# ==============================================================================
egyptian_stocks = [
    "AALR.CA", "ABUK.CA", "ACAMD.CA", "ACAP.CA", "ACGC.CA", "ACTF.CA", "ADCI.CA", "ADIB.CA",
    "ADPC.CA", "ADRI.CA", "AFDI.CA", "AFMC.CA", "AIDC.CA", "AIFI.CA", "AIH.CA", "AJWA.CA",
    "ALCN.CA", "ALEX.CA", "ALUM.CA", "AMER.CA", "AMES.CA", "AMIA.CA", "AMII.CA", "AMOC.CA",
    "AMPI.CA", "APSW.CA", "ARAB.CA", "ARCC.CA", "AREH.CA", "ASCM.CA", "ASPI.CA", "ATLC.CA",
    "ATQA.CA", "AXPH.CA", "BIDI.CA", "BIGP.CA", "BINV.CA", "BIOC.CA", "BONY.CA", "BTFH.CA",
    "CAED.CA", "CANA.CA", "CCAP.CA", "CCRS.CA", "CEFM.CA", "CERA.CA", "CFGH.CA", "CICH.CA",
    "CIEB.CA", "CIRA.CA", "CLHO.CA", "CNFN.CA", "COMI.CA", "COPR.CA", "COSG.CA", "CPCI.CA",
    "CPME.CA", "CRST.CA", "CSAG.CA", "DAPH.CA", "DCRC.CA", "DEIN.CA", "DGTZ.CA", "DOMT.CA",
    "DSCW.CA", "DTPP.CA", "EALR.CA", "EASB.CA", "EAST.CA", "EBSC.CA", "ECAP.CA", "EDFM.CA",
    "EEII.CA", "EFIC.CA", "EFID.CA", "EFIH.CA", "EGAL.CA", "EGAS.CA", "EGBE.CA", "EGCH.CA",
    "EGREF.CA", "EGSA.CA", "EGTS.CA", "EHDR.CA", "ELAB.CA", "ELEC.CA", "ELKA.CA", "ELNA.CA",
    "ELSH.CA", "ELWA.CA", "EMFD.CA", "ENGC.CA", "EOSB.CA", "EPCO.CA", "EPPK.CA", "ETEL.CA",
    "ETRS.CA", "EXPA.CA", "FAIT.CA", "FAITA.CA", "FCMD.CA", "FIRE.CA", "FNAR.CA", "FTNS.CA",
    "FWRY.CA", "GBCO.CA", "GDWA.CA", "GGCC.CA", "GGRN.CA", "GIHD.CA", "GMCI.CA", "GOUR.CA",
    "GPIM.CA", "GRCA.CA", "GSSC.CA", "GTEX.CA", "GTHE.CA", "GTWL.CA", "HBCO.CA", "HDBK.CA",
    "HELI.CA", "HRHO.CA", "IBCT.CA", "ICFC.CA", "ICID.CA", "IDRE.CA", "IEEC.CA", "IFAP.CA",
    "INEG.CA", "INFI.CA", "IRON.CA", "ISMA.CA", "ISMQ.CA", "ISPH.CA", "JUFO.CA", "KABO.CA",
    "KORA.CA", "KRDI.CA", "KWIN.CA", "KZPC.CA", "LCSW.CA", "LKGP.CA", "LUTS.CA", "MAAL.CA",
    "MASR.CA", "MBEG.CA", "MBSC.CA", "MCQE.CA", "MCRO.CA", "MENA.CA", "MEPA.CA", "MFPC.CA",
    "MFSC.CA", "MHOT.CA", "MILS.CA", "MIPH.CA", "MOED.CA", "MOIL.CA", "MOIN.CA", "MOSC.CA",
    "MPCI.CA", "MPCO.CA", "MPRC.CA", "MTIE.CA", "NAHO.CA", "NARE.CA", "NCCW.CA", "NCGC.CA",
    "NEDA.CA", "NHPS.CA", "NINH.CA", "NIPH.CA", "OBRI.CA", "OCAP.CA", "OCDI.CA", "OCPH.CA",
    "ODIN.CA", "OFH.CA", "OIH.CA", "OLFI.CA", "ORAS.CA", "ORHD.CA", "ORWE.CA", "PHAR.CA",
    "PHDC.CA", "PHGC.CA", "PHTV.CA", "POUL.CA", "PRCL.CA", "PRDC.CA", "PRMH.CA", "QNBE.CA",
    "RACC.CA", "RAKT.CA", "RAYA.CA", "RKAZ.CA", "RMDA.CA", "RMTV.CA", "ROTO.CA", "RREI.CA",
    "RTVC.CA", "RUBX.CA", "SAUD.CA", "SCEM.CA", "SCFM.CA", "SCTS.CA", "SDTI.CA", "SEIG.CA",
    "SIEG.CA", "SIPC.CA", "SKPC.CA", "SMFR.CA", "SNFC.CA", "SPIN.CA", "SPMD.CA", "SUCE.CA",
    "SUGR.CA", "SVCE.CA", "SWDY.CA", "TALM.CA", "TANM.CA", "TAQA.CA", "TMGH.CA", "TORA.CA",
    "TWSA.CA", "TYCN.CA", "UBEE.CA", "UEFM.CA", "UEGC.CA", "UNIP.CA", "UNIT.CA", "UPMS.CA",
    "UTOP.CA", "VALU.CA", "VERT.CA", "VLMR.CA", "VLMRA.CA", "WCDF.CA", "WKOL.CA", "ZEOT.CA",
    "ZMID.CA"
]

# ==============================================================================
# 2. الهياكل الخوارزمية لاكتشاف نموذج Adam & Adam
# ==============================================================================
class PatternStatus(Enum):
    FORMING = "forming"
    CONFIRMED = "confirmed"
    FAILED = "failed"
    COMPLETED = "completed"

class BottomType(Enum):
    ADAM = "adam"
    EVE = "eve"
    UNKNOWN = "unknown"

DEFAULT_PARAMS = {
    'min_peak_rise_pct': 10.0,
    'max_bottom_price_diff_pct': 3.0,
    'min_separation_bars': 15,
    'max_separation_bars': 45,
    'adam_width_threshold': 0.4,
    'min_prior_downtrend_pct': 10.0,
    'prior_trend_lookback': 40,
    'order': 3,
}

@dataclass
class Bottom:
    index: int
    date: pd.Timestamp
    price: float
    volume: float
    width: int = 0
    spike_intensity: float = 0.0
    bottom_type: BottomType = BottomType.UNKNOWN

@dataclass    
class AdamAdamDoubleBottom:
    first_bottom: Bottom
    second_bottom: Bottom
    peak_index: int
    peak_date: pd.Timestamp
    peak_price: float
    status: PatternStatus = PatternStatus.FORMING
    pattern_height: float = 0.0
    pattern_width: int = 0
    price_difference_pct: float = 0.0
    peak_rise_pct: float = 0.0
    prior_downtrend_pct: float = 0.0
    confirmation_index: Optional[int] = None
    confirmation_date: Optional[pd.Timestamp] = None
    entry_price: Optional[float] = None
    stop_loss: Optional[float] = None
    price_target: Optional[float] = None

    def __post_init__(self):
        self.pattern_height = self.peak_price - min(self.first_bottom.price, self.second_bottom.price)
        self.pattern_width = self.second_bottom.index - self.first_bottom.index
        self.price_difference_pct = abs(self.first_bottom.price - self.second_bottom.price) / self.first_bottom.price * 100
        self.peak_rise_pct = (self.peak_price - self.first_bottom.price) / self.first_bottom.price * 100

class AdamAdamDetector:
    def __init__(self, params: Optional[Dict] = None):
        self.params = {**DEFAULT_PARAMS, **(params or {})}
        self.patterns: List[AdamAdamDoubleBottom] = []
        self._data: Optional[pd.DataFrame] = None

    def detect(self, data: pd.DataFrame) -> List[AdamAdamDoubleBottom]:
        if len(data) < (self.params['min_separation_bars'] + self.params['prior_trend_lookback']):
            return []
        self._data = data.copy()
        self.patterns = []

        bottoms = self._find_all_bottoms()
        classified_bottoms = self._classify_bottoms(bottoms)
        candidate_pairs = self._find_adam_adam_pairs(classified_bottoms)

        for first, second, peak_idx, prior_drop_pct in candidate_pairs:
            pattern = self._create_pattern(first, second, peak_idx, prior_drop_pct)
            if pattern:
                self.patterns.append(pattern)

        self._check_confirmations()
        return self.patterns

    def _find_all_bottoms(self) -> List[Bottom]:
        closes = self._data['Close'].values
        volumes = self._data['Volume'].values if 'Volume' in self._data.columns else np.ones(len(closes))
        minima_indices = argrelextrema(closes, np.less, order=self.params['order'])[0]

        bottoms = []
        for idx in minima_indices:
            bottoms.append(Bottom(
                index=idx,
                date=self._data.index[idx],
                price=self._data['Low'].iloc[idx],
                volume=volumes[idx] if idx < len(volumes) else 0
            ))
        return bottoms

    def _classify_bottoms(self, bottoms: List[Bottom]) -> List[Bottom]:
        closes = self._data['Close'].values
        lows = self._data['Low'].values

        for bottom in bottoms:
            idx = bottom.index
            window = max(5, self.params['order'] * 2)

            pre_bottom_level = closes[max(0, idx - window)]
            recovery_bars = 0
            for i in range(idx + 1, min(len(closes), idx + window * 2)):
                if closes[i] >= pre_bottom_level:
                    recovery_bars = i - idx
                    break
            bottom.width = recovery_bars or window

            decline_window = min(5, idx)
            if decline_window > 1:
                pre_decline = closes[max(0, idx - decline_window)]
                spike_pct = (pre_decline - lows[idx]) / pre_decline * 100
                bottom.spike_intensity = spike_pct / max(1, decline_window)

            width_score = bottom.width / window
            spike_score = min(bottom.spike_intensity / 3.0, 1.0)

            if width_score < self.params['adam_width_threshold'] and spike_score > 0.3:
                bottom.bottom_type = BottomType.ADAM
            else:
                bottom.bottom_type = BottomType.EVE

        return bottoms

    def _find_adam_adam_pairs(self, bottoms: List[Bottom]) -> List[Tuple[Bottom, Bottom, int, float]]:
        adam_bottoms = [b for b in bottoms if b.bottom_type == BottomType.ADAM]
        pairs = []

        for i in range(len(adam_bottoms)):
            first = adam_bottoms[i]

            lookback_start = max(0, first.index - self.params['prior_trend_lookback'])
            if first.index - lookback_start < 10:
                continue

            prior_high = self._data['High'].iloc[lookback_start:first.index].max()
            prior_drop_pct = (prior_high - first.price) / prior_high * 100

            if prior_drop_pct < self.params['min_prior_downtrend_pct']:
                continue

            for j in range(i + 1, len(adam_bottoms)):
                second = adam_bottoms[j]
                separation = second.index - first.index
                if not (self.params['min_separation_bars'] <= separation <= self.params['max_separation_bars']):
                    continue

                if abs(first.price - second.price) / first.price * 100 > self.params['max_bottom_price_diff_pct']:
                    continue

                peak_idx = self._find_peak_between(first.index, second.index)
                if peak_idx is None:
                    continue

                peak_price = self._data['High'].iloc[peak_idx]
                if (peak_price - first.price) / first.price * 100 < self.params['min_peak_rise_pct']:
                    continue

                pairs.append((first, second, peak_idx, prior_drop_pct))
        return pairs

    def _find_peak_between(self, start_idx: int, end_idx: int) -> Optional[int]:
        if end_idx <= start_idx + 1:
            return None
        highs = self._data['High'].iloc[start_idx:end_idx]
        peak_idx = highs.idxmax()
        peak_pos = self._data.index.get_loc(peak_idx)
        return peak_pos if start_idx < peak_pos < end_idx else None

    def _create_pattern(self, first: Bottom, second: Bottom, peak_idx: int, prior_drop_pct: float) -> Optional[AdamAdamDoubleBottom]:
        peak_price = self._data['High'].iloc[peak_idx]
        pattern = AdamAdamDoubleBottom(
            first_bottom=first,
            second_bottom=second,
            peak_index=peak_idx,
            peak_date=self._data.index[peak_idx],
            peak_price=peak_price,
            prior_downtrend_pct=prior_drop_pct
        )
        pattern.entry_price = pattern.peak_price * 1.02
        pattern.stop_loss = pattern.peak_price * 0.94
        pattern.price_target = pattern.peak_price + pattern.pattern_height * 0.73
        return pattern

    def _check_confirmations(self):
        for pattern in self.patterns:
            second_idx = pattern.second_bottom.index
            subsequent_data = self._data.iloc[second_idx + 1:]

            for idx, row in subsequent_data.iterrows():
                current_idx = self._data.index.get_loc(idx)
                if row['Close'] > pattern.peak_price:
                    pattern.status = PatternStatus.CONFIRMED
                    pattern.confirmation_index = current_idx
                    pattern.confirmation_date = idx
                    break
                if row['Close'] < min(pattern.first_bottom.price, pattern.second_bottom.price) * 0.98:
                    pattern.status = PatternStatus.FAILED
                    break

# ==============================================================================
# 3. دمج بيانات TradingView اللحظية مع Yahoo Finance لآخر سنة
# ==============================================================================
def get_combined_stock_data():
    end_date = datetime.datetime.now()
    start_date = end_date - datetime.timedelta(days=365)

    print("1️⃣ جلب أحدث البيانات اللحظية من TradingView...")
    tv_url = "https://scanner.tradingview.com/egypt/scan"
    tv_payload = {
        "filter": [{"left": "name", "operation": "nempty"}],
        "options": {"active_symbols_only": True},
        "columns": ["name", "description", "open", "high", "low", "close"],
        "sort": {"sortBy": "name", "sortOrder": "asc"},
        "range": [0, 300],
    }
    tv_headers = {"User-Agent": "Mozilla/5.0"}
    
    tv_dict = {}
    today_date = pd.Timestamp(datetime.date.today())
    try:
        res = requests.post(tv_url, json=tv_payload, headers=tv_headers, timeout=10)
        data = res.json().get("data", [])
        for item in data:
            sym = f"{item['s'].replace('EGX:', '')}.CA"
            d = item["d"]
            open_p = float(d[2]) if d[2] is not None else 0.0
            high_p = float(d[3]) if d[3] is not None else open_p
            low_p = float(d[4]) if d[4] is not None else open_p
            close_p = float(d[5]) if d[5] is not None else open_p
            
            if close_p > 0:
                tv_dict[sym] = {
                    'Open': open_p,
                    'High': high_p,
                    'Low': low_p,
                    'Close': close_p,
                    'Volume': 0
                }
    except Exception as e:
        print(f"⚠️ تعذر جلب بيانات TradingView: {e}")

    print("2️⃣ جلب بيانات آخر سنة من Yahoo Finance مع إخفاء الأخطاء...")
    stderr_backup = sys.stderr
    sys.stderr = open(os.devnull, "w")
    try:
        bulk_data = yf.download(egyptian_stocks, start=start_date, end=end_date, group_by='ticker', threads=True)
    finally:
        sys.stderr = stderr_backup

    combined_data = {}
    print("3️⃣ دمج وسد الفجوات بين المصدرين حتى اللحظة الحالية...")
    for ticker in egyptian_stocks:
        try:
            if ticker in bulk_data.columns.levels[0]:
                df_stock = bulk_data[ticker].dropna(how='all').copy()
                
                if ticker in tv_dict:
                    last_tv_data = tv_dict[ticker]
                    if df_stock.empty or df_stock.index[-1].date() < today_date.date():
                        df_tv_row = pd.DataFrame([last_tv_data], index=[today_date])
                        df_stock = pd.concat([df_stock, df_tv_row])
                
                if not df_stock.empty:
                    df_stock = df_stock.ffill().bfill()
                    combined_data[ticker] = df_stock
        except Exception:
            continue

    return combined_data

# ==============================================================================
# 4. المعالجة المتوازية ورصد الصفقات (المغلقة والمفتوحة) مع تحديث الـ State
# ==============================================================================
all_stock_datasets = get_combined_stock_data()

def process_single_stock(ticker, hold_days=60):
    if ticker not in all_stock_datasets:
        return []

    df_stock = all_stock_datasets[ticker]
    if len(df_stock) < 35:
        return []

    detector = AdamAdamDetector()
    patterns = detector.detect(df_stock)

    trades = []
    latest_price = df_stock['Close'].iloc[-1]

    for pattern in patterns:
        if pattern.status != PatternStatus.CONFIRMED:
            continue

        entry_idx = pattern.confirmation_index
        if entry_idx is None or entry_idx >= len(df_stock):
            continue

        entry_price = pattern.entry_price or (df_stock['Close'].iloc[entry_idx] * 1.02)
        stop_price = pattern.stop_loss or (pattern.peak_price * 0.94)
        target_price = pattern.price_target or entry_price * 1.2

        exit_date = None
        state = 'Open'

        # فحص حركة السهم بعد تاريخ الدخول
        for i in range(1, len(df_stock) - entry_idx):
            current_idx = entry_idx + i
            current_high = df_stock['High'].iloc[current_idx]
            current_low = df_stock['Low'].iloc[current_idx]

            if current_low <= stop_price:
                exit_date = df_stock.index[current_idx]
                state = 'Loss'
                break

            if current_high >= target_price:
                exit_date = df_stock.index[current_idx]
                state = 'Win'
                break
            
            # إنهاء الصفقات بعد تجاوز فترة الاحتفاظ الإجمالية
            if i >= hold_days:
                exit_date = df_stock.index[current_idx]
                close_p = df_stock['Close'].iloc[current_idx]
                state = 'Win' if close_p >= entry_price else 'Loss'
                break

        # تاريخ الخروج يترك فارغاً إذا كانت الصفقة Open
        formatted_exit_date = exit_date.strftime('%Y-%m-%d') if (state != 'Open' and hasattr(exit_date, 'strftime') and exit_date) else ''

        # تنسيق الأعمدة الثمانية المطلوبة باللغة الإنجليزية
        trades.append({
            'Stock Name': ticker,
            'Entry Date': pattern.confirmation_date.strftime('%Y-%m-%d') if pattern.confirmation_date else '',
            'Entry Price': round(entry_price, 2),
            'Target': round(target_price, 2),
            'Stop Loss': round(stop_price, 2),
            'Exit Date': formatted_exit_date,
            'Current Price': round(latest_price, 2),
            'State': state
        })

    return trades

print("⚡ جارٍ الفحص لآخر سنة ورصد الصفقات المفتوحة والمغلقة...")
all_trades = []
with concurrent.futures.ThreadPoolExecutor(max_workers=16) as executor:
    results = list(executor.map(process_single_stock, egyptian_stocks))
    for res in results:
        all_trades.extend(res)

# ==============================================================================
# 5. التنسيق وإصدار ملف Excel المفصل بالتصميم والألوان المحددة
# ==============================================================================
df_trades = pd.DataFrame(all_trades)

def format_excel_file(excel_filename: str):
    wb = openpyxl.load_workbook(excel_filename)
    ws = wb.active
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1B4266", end_color="1B4266", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    data_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # التنسيق الشرطي لحالة الصفقة (State)
    win_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    win_font = Font(name="Segoe UI", size=10, bold=True, color="155724")

    loss_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    loss_font = Font(name="Segoe UI", size=10, bold=True, color="721C24")

    open_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    open_font = Font(name="Segoe UI", size=10, bold=True, color="856404")

    ws.row_dimensions[1].height = 26
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 20
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

            # تطبيق التلوين الذكي لعمود State
            if col == 8:
                val = str(cell.value).strip()
                if val == 'Win':
                    cell.fill = win_fill
                    cell.font = win_font
                elif val == 'Loss':
                    cell.fill = loss_fill
                    cell.font = loss_font
                elif val == 'Open':
                    cell.fill = open_fill
                    cell.font = open_font

    column_widths = {'A': 18, 'B': 16, 'C': 16, 'D': 16, 'E': 16, 'F': 16, 'G': 16, 'H': 14}
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(excel_filename)

excel_filename = "Adam_Adam_Pattern_Single_Sheet.xlsx"

if not df_trades.empty:
    # 🛑 1. منع تكرار الصفقات بنفس السهم وتاريخ الدخول
    df_trades = df_trades.drop_duplicates(subset=['Stock Name', 'Entry Date'], keep='first')

    # 📅 2. الترتيب الزمني التلقائي من الأقدم إلى الأحدث
    df_trades['Entry Date Temp'] = pd.to_datetime(df_trades['Entry Date'])
    df_trades = df_trades.sort_values(by='Entry Date Temp', ascending=True).drop(columns=['Entry Date Temp'])

    # 3. التأكد من ترتيب الأعمدة المطلوب
    required_columns = [
        'Stock Name', 'Entry Date', 'Entry Price', 'Target', 
        'Stop Loss', 'Exit Date', 'Current Price', 'State'
    ]
    df_trades = df_trades[required_columns]

    try:
        with pd.ExcelWriter(excel_filename, engine='openpyxl') as writer:
            df_trades.to_excel(writer, sheet_name='Trades List', index=False)

        format_excel_file(excel_filename)
        print(f"\n📊 تم إنشاء الملف بنجاح وتحديث الصفقات لآخر سنة: {excel_filename}")
        print(f"إجمالي الصفقات: {len(df_trades)} (تتضمن الصفقات المغلقة والصفقات المفتوحة حالياً)")
    except PermissionError:
        print(f"\n❌ يرجى إغلاق ملف '{excel_filename}' من Excel أولاً ثم إعادة التشغيل!")
else:
    df_empty = pd.DataFrame(columns=[
        'Stock Name', 'Entry Date', 'Entry Price', 'Target', 
        'Stop Loss', 'Exit Date', 'Current Price', 'State'
    ])
    with pd.ExcelWriter(excel_filename, engine='openpyxl') as writer:
        df_empty.to_excel(writer, sheet_name='Trades List', index=False)
    format_excel_file(excel_filename)
    print("\n⚠️ لم يتم العثور على أي صفقات تطابق الشروط خلال آخر سنة، وتم إنشاء الهيكل الفارغ بنجاح.")
