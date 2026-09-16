"""
================================================================================
EVE & EVE BACKTESTER & SCANNER (10-YEAR DATA + PERFORMANCE METRICS IN TERMINAL)
================================================================================
"""

import math
import warnings
from dataclasses import dataclass
from datetime import datetime, timedelta
from enum import Enum
from typing import List, Optional, Tuple

import numpy as np
import openpyxl
import pandas as pd
import requests
import yfinance as yf
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from scipy.signal import argrelextrema

warnings.filterwarnings('ignore')

try:
    from google.colab import files
    COLAB_ENV = True
except ImportError:
    COLAB_ENV = False


# ====================================================================
# 1. TRADINGVIEW INTEGRATION FOR LIVE DATA
# ====================================================================

def fetch_tradingview_live_data(tickers: List[str]) -> dict:
    """Fetches real-time price data from TradingView scanner for EGX stocks."""
    tv_data = {}
    symbols = [f"EGX:{ticker.replace('.CA', '')}" for ticker in tickers]
    
    url = "https://scanner.tradingview.com/egypt/scan"
    payload = {
        "symbols": {"tickers": symbols},
        "columns": ["name", "open", "high", "low", "close", "volume"]
    }
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
    }

    try:
        response = requests.post(url, json=payload, headers=headers, timeout=10)
        if response.status_code == 200:
            res_json = response.json()
            for item in res_json.get("data", []):
                raw_ticker = item.get("s", "")
                ticker = raw_ticker.replace("EGX:", "") + ".CA"
                vals = item.get("d", [])
                if len(vals) >= 6 and vals[1] is not None:
                    tv_data[ticker] = {
                        "Open": float(vals[1]),
                        "High": float(vals[2]),
                        "Low": float(vals[3]),
                        "Close": float(vals[4]),
                        "Volume": float(vals[5]) if vals[5] is not None else 0.0
                    }
    except Exception as e:
        print(f"    [TV Warning] Could not fetch live data from TradingView: {e}")

    return tv_data


def get_combined_stock_data(ticker: str, start_date: datetime, end_date: datetime, tv_live_dict: dict) -> pd.DataFrame:
    """Downloads Yahoo Finance data and merges TradingView live data for missing current bar."""
    df = yf.download(
        ticker,
        start=start_date.strftime('%Y-%m-%d'),
        end=end_date.strftime('%Y-%m-%d'),
        progress=False
    )

    if df.empty:
        return pd.DataFrame()

    if isinstance(df.columns, pd.MultiIndex):
        df.columns = df.columns.get_level_values(0)

    # Attach TV Live data if today's candle is not updated in YFinance
    if ticker in tv_live_dict:
        tv_bar = tv_live_dict[ticker]
        today_date = pd.Timestamp(datetime.now().date())
        
        if df.index[-1].date() < today_date.date():
            new_row = pd.DataFrame({
                'Open': [tv_bar['Open']],
                'High': [tv_bar['High']],
                'Low': [tv_bar['Low']],
                'Close': [tv_bar['Close']],
                'Volume': [tv_bar['Volume']]
            }, index=[today_date])
            df = pd.concat([df, new_row])
        else:
            df.iloc[-1, df.columns.get_loc('High')] = max(df.iloc[-1]['High'], tv_bar['High'])
            df.iloc[-1, df.columns.get_loc('Low')] = min(df.iloc[-1]['Low'], tv_bar['Low'])
            df.iloc[-1, df.columns.get_loc('Close')] = tv_bar['Close']
            df.iloc[-1, df.columns.get_loc('Volume')] = max(df.iloc[-1]['Volume'], tv_bar['Volume'])

    df.ffill(inplace=True)
    df.bfill(inplace=True)
    return df


# ====================================================================
# 2. PATTERN DETECTOR CLASSES
# ====================================================================

class BottomType(Enum):
    ADAM = 'Adam'
    EVE = 'Eve'
    UNKNOWN = 'Unknown'


@dataclass
class Bottom:
    index: int
    date: pd.Timestamp
    price: float
    bottom_type: BottomType
    width: int
    depth_pct: float
    volume_ratio: float
    sharpness: float
    roundness: float


@dataclass
class EveEvePattern:
    first_bottom: Bottom   # Eve (القاع المدور أولاً)
    second_bottom: Bottom  # Eve (القاع المدور ثانياً)
    peak: float
    peak_index: int
    peak_date: pd.Timestamp
    separation_days: int
    depth_pct: float
    rise_between_pct: float
    valid: bool
    score: float
    breakout_index: Optional[int] = None
    breakout_date: Optional[pd.Timestamp] = None
    breakout_price: Optional[float] = None
    breakout_volume_ratio: Optional[float] = None
    entry_price: Optional[float] = None
    target_price: Optional[float] = None
    stop_loss: Optional[float] = None


class EveEveDetector:

    def __init__(
        self,
        min_trend_bars: int = 15,
        min_trend_decline_pct: float = 8.0,
        order: int = 3,
        adam_max_width: int = 4,
        adam_min_sharpness: float = 0.3,
        adam_min_volume_ratio: float = 1.0,
        eve_min_width: int = 4,
        eve_min_roundness: float = 0.3,
        eve_max_sharpness: float = 0.5,
        min_separation_days: int = 15,
        max_separation_days: int = 56,
        min_rise_between_pct: float = 10.0,
        bottom_similarity_pct: float = 4.0,
        breakout_buffer_pct: float = 1.5,
        stop_loss_pct_below_peak: float = 5.0,
        min_depth_pct: float = 6.0,
        max_depth_pct: float = 60.0,
        breakout_volume_ratio: float = 0.8,
        max_bars_after_second_bottom: int = 60,
    ):
        self.min_trend_bars = min_trend_bars
        self.min_trend_decline_pct = min_trend_decline_pct
        self.order = order
        self.adam_max_width = adam_max_width
        self.adam_min_sharpness = adam_min_sharpness
        self.adam_min_volume_ratio = adam_min_volume_ratio
        self.eve_min_width = eve_min_width
        self.eve_min_roundness = eve_min_roundness
        self.eve_max_sharpness = eve_max_sharpness

        self.min_separation_days = min_separation_days
        self.max_separation_days = max_separation_days
        self.min_rise_between_pct = min_rise_between_pct
        self.bottom_similarity_pct = bottom_similarity_pct
        self.breakout_buffer_pct = breakout_buffer_pct
        self.stop_loss_pct_below_peak = stop_loss_pct_below_peak

        self.min_depth_pct = min_depth_pct
        self.max_depth_pct = max_depth_pct
        self.breakout_volume_ratio = breakout_volume_ratio
        self.max_bars_after_second_bottom = max_bars_after_second_bottom

    def detect(self, df: pd.DataFrame) -> List[EveEvePattern]:
        df = self._prepare_dataframe(df)
        bottoms = self._find_local_minima(df)
        classified = self._classify_bottoms(df, bottoms)

        eves = [b for b in classified if b.bottom_type == BottomType.EVE]

        patterns = self._match_eve_eve(df, eves)
        patterns = self._check_breakouts(df, patterns)
        patterns = self._score_patterns(patterns)
        patterns = self._enforce_exclusivity(patterns)

        patterns.sort(key=lambda x: x.score, reverse=True)
        return patterns

    def _prepare_dataframe(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        if not isinstance(df.index, pd.DatetimeIndex):
            df.index = pd.to_datetime(df.index)
        required_cols = ['Open', 'High', 'Low', 'Close', 'Volume']
        for col in required_cols:
            if col not in df.columns:
                raise ValueError(f"Required column '{col}' not found")
        return df

    def _find_local_minima(self, df: pd.DataFrame) -> List[Tuple[int, float]]:
        lows = df['Low'].values
        minima_idx = argrelextrema(lows, np.less, order=self.order)[0]
        bottoms = []
        for idx in minima_idx:
            ws = max(0, idx - self.order)
            we = min(len(df), idx + self.order + 1)
            if lows[idx] == lows[ws:we].min():
                bottoms.append((idx, lows[idx]))
        return bottoms

    def _classify_bottoms(self, df: pd.DataFrame, bottoms: List[Tuple[int, float]]) -> List[Bottom]:
        classified = []
        highs = df['High'].values
        lows = df['Low'].values
        volumes = df['Volume'].values
        avg_volume = volumes.mean()

        for idx, price in bottoms:
            left = max(0, idx - self.order * 3)
            right = min(len(df), idx + self.order * 3 + 1)
            window_lows = lows[left:right]
            window_highs = highs[left:right]

            threshold = price * 1.03
            width = int((window_lows <= threshold).sum())

            surrounding_high = max(
                window_highs.max(),
                highs[max(0, idx - 1)],
                highs[min(len(df) - 1, idx + 1)],
            )
            depth_pct = ((surrounding_high - price) / surrounding_high) * 100

            volume_ratio = volumes[idx] / avg_volume if avg_volume > 0 else 1.0
            sharpness = self._calculate_sharpness(df, idx)
            roundness = self._calculate_roundness(df, idx)

            adam_score = 0
            eve_score = 0

            if width <= self.adam_max_width:
                adam_score += 2
            if sharpness >= self.adam_min_sharpness:
                adam_score += 2
            if volume_ratio >= self.adam_min_volume_ratio:
                adam_score += 1
            if roundness < self.eve_min_roundness:
                adam_score += 1

            if width >= self.eve_min_width:
                eve_score += 2
            if roundness >= self.eve_min_roundness:
                eve_score += 2
            if sharpness <= self.eve_max_sharpness:
                eve_score += 1

            if adam_score >= 4 and eve_score < 3:
                btype = BottomType.ADAM
            elif eve_score >= 4 and adam_score < 3:
                btype = BottomType.EVE
            elif adam_score > eve_score + 1:
                btype = BottomType.ADAM
            elif eve_score > adam_score + 1:
                btype = BottomType.EVE
            else:
                btype = BottomType.UNKNOWN

            classified.append(
                Bottom(
                    index=idx,
                    date=df.index[idx],
                    price=price,
                    bottom_type=btype,
                    width=width,
                    depth_pct=depth_pct,
                    volume_ratio=volume_ratio,
                    sharpness=sharpness,
                    roundness=roundness,
                )
            )
        return classified

    def _calculate_sharpness(self, df: pd.DataFrame, idx: int) -> float:
        lows = df['Low'].values
        highs = df['High'].values
        closes = df['Close'].values

        pre = max(0, idx - 4)
        post = min(len(df), idx + 5)
        if idx - pre < 2 or post - idx < 2:
            return 0.0

        pre_high = highs[pre:idx].max()
        post_high = highs[idx + 1 : post].max()

        decline = (pre_high - lows[idx]) / max(idx - pre, 1)
        recovery = (post_high - lows[idx]) / max(post - idx - 1, 1)
        sharpness = min(decline, recovery) / (pre_high * 0.005 + 1e-10)

        if idx > 0 and idx < len(df) - 1:
            if (
                closes[idx - 1] > closes[idx]
                and closes[idx + 1] > closes[idx]
                and lows[idx] < lows[idx - 1]
                and lows[idx] < lows[idx + 1]
            ):
                sharpness *= 1.5

        return min(sharpness, 3.0)

    def _calculate_roundness(self, df: pd.DataFrame, idx: int) -> float:
        lows = df['Low'].values
        left = max(0, idx - self.order * 3)
        right = min(len(df), idx + self.order * 3 + 1)
        window = lows[left:right]
        if len(window) < 5:
            return 0.0

        x = np.arange(len(window))
        y_norm = (window - window.min()) / (window.max() - window.min() + 1e-10)

        try:
            coeffs = np.polyfit(x, y_norm, 2)
        except Exception:
            return 0.0

        if coeffs[0] <= 0:
            return 0.0

        y_pred = np.polyval(coeffs, x)
        ss_res = np.sum((y_norm - y_pred) ** 2)
        ss_tot = np.sum((y_norm - y_norm.mean()) ** 2)
        r_squared = 1 - (ss_res / (ss_tot + 1e-10))

        roundness = r_squared * min(coeffs[0] * 8, 1.0)
        low_range = window.max() - window.min()
        flatness = 1 - min(low_range / (window.mean() + 1e-10), 1.0)
        roundness += flatness * 0.3

        return min(roundness, 2.0)

    def _match_eve_eve(self, df: pd.DataFrame, eves: List[Bottom]) -> List[EveEvePattern]:
        """Matches Eve (first) with Eve (second) bottoms."""
        patterns = []
        highs = df['High'].values

        for i in range(len(eves)):
            for j in range(i + 1, len(eves)):
                first = eves[i]
                second = eves[j]

                separation = second.index - first.index
                if not (self.min_separation_days <= separation <= self.max_separation_days):
                    continue

                between_highs = highs[first.index : second.index]
                if len(between_highs) == 0:
                    continue

                peak_idx_rel = between_highs.argmax()
                peak_idx = first.index + peak_idx_rel
                peak_price = between_highs[peak_idx_rel]

                if peak_idx_rel == 0 or peak_idx_rel == len(between_highs) - 1:
                    continue

                bottom_diff = abs(second.price - first.price) / first.price * 100
                if bottom_diff > self.bottom_similarity_pct:
                    continue

                rise_from_first_pct = ((peak_price - first.price) / first.price) * 100
                rise_from_second_pct = ((peak_price - second.price) / second.price) * 100

                if (rise_from_first_pct < self.min_rise_between_pct or rise_from_second_pct < self.min_rise_between_pct):
                    continue

                avg_bottom = (first.price + second.price) / 2
                depth_pct = ((peak_price - avg_bottom) / peak_price) * 100

                if not self._check_preceding_trend(df, first.index):
                    continue

                patterns.append(
                    EveEvePattern(
                        first_bottom=first,
                        second_bottom=second,
                        peak=peak_price,
                        peak_index=peak_idx,
                        peak_date=df.index[peak_idx],
                        separation_days=separation,
                        depth_pct=depth_pct,
                        rise_between_pct=rise_from_first_pct,
                        valid=True,
                        score=0.0,
                    )
                )

        return patterns

    def _check_preceding_trend(self, df: pd.DataFrame, eve_idx: int) -> bool:
        if eve_idx < self.min_trend_bars:
            return False
        start_idx = eve_idx - self.min_trend_bars
        trend_high = df['High'].iloc[start_idx:eve_idx].max()
        eve_low = df['Low'].iloc[eve_idx]
        decline_pct = ((trend_high - eve_low) / trend_high) * 100
        return decline_pct >= self.min_trend_decline_pct

    def _check_breakouts(self, df: pd.DataFrame, patterns: List[EveEvePattern]) -> List[EveEvePattern]:
        closes = df['Close'].values
        volumes = df['Volume'].values
        avg_volume = volumes.mean()

        for p in patterns:
            second_idx = p.second_bottom.index
            max_check = min(len(df), second_idx + self.max_bars_after_second_bottom + 1)
            required_breakout_price = p.peak * (1 + (self.breakout_buffer_pct / 100.0))

            breakout_found = False
            for i in range(second_idx + 1, max_check):
                if closes[i] >= required_breakout_price:
                    vol_ratio = volumes[i] / avg_volume if avg_volume > 0 else 0
                    if vol_ratio >= self.breakout_volume_ratio:
                        p.breakout_index = i
                        p.breakout_date = df.index[i]
                        p.breakout_price = closes[i]
                        p.breakout_volume_ratio = vol_ratio
                        p.entry_price = required_breakout_price

                        height = p.peak - min(p.first_bottom.price, p.second_bottom.price)
                        p.target_price = p.peak + height
                        p.stop_loss = p.peak * (1 - (self.stop_loss_pct_below_peak / 100.0))

                        breakout_found = True
                        break

            if not breakout_found:
                p.valid = False

        return [p for p in patterns if p.valid]

    def _score_patterns(self, patterns: List[EveEvePattern]) -> List[EveEvePattern]:
        for p in patterns:
            score = 50.0

            if 21 <= p.separation_days <= 35:
                score += 15
            elif 36 <= p.separation_days <= 56:
                score += 10

            if 15 <= p.rise_between_pct <= 30:
                score += 15
            elif p.rise_between_pct > 30:
                score += 10

            bottom_diff = (
                abs(p.first_bottom.price - p.second_bottom.price)
                / p.first_bottom.price
                * 100
            )
            if bottom_diff < 1.0:
                score += 15
            elif bottom_diff <= 3.0:
                score += 10

            if p.second_bottom.volume_ratio > p.first_bottom.volume_ratio:
                score += 10

            p.score = min(score, 100)

        return patterns

    def _enforce_exclusivity(self, patterns: List[EveEvePattern]) -> List[EveEvePattern]:
        if not patterns:
            return []

        patterns.sort(key=lambda x: x.score, reverse=True)

        filtered = []
        used_bottom_pairs = set()
        used_indices = set()

        for p in patterns:
            end_idx = p.breakout_index if p.breakout_index else p.second_bottom.index + 5
            p_range = set(range(p.first_bottom.index - 3, end_idx + 3))
            bottom_pair = (p.first_bottom.index, p.second_bottom.index)

            has_overlap = bool(p_range.intersection(used_indices))
            bottoms_used = bottom_pair in used_bottom_pairs

            if not has_overlap and not bottoms_used:
                filtered.append(p)
                used_indices.update(p_range)
                used_bottom_pairs.add(bottom_pair)

        return filtered


# ====================================================================
# 3. LIST OF EGYPTIAN STOCKS
# ====================================================================

EGYPTIAN_STOCKS = [
    'AALR.CA', 'ABUK.CA', 'ACAMD.CA', 'ACAP.CA', 'ACGC.CA', 'ACTF.CA', 'ADCI.CA',
    'ADIB.CA', 'ADPC.CA', 'ADRI.CA', 'AFDI.CA', 'AFMC.CA', 'AIDC.CA', 'AIFI.CA',
    'AIH.CA', 'AJWA.CA', 'ALCN.CA', 'ALEX.CA', 'ALUM.CA', 'AMER.CA', 'AMES.CA',
    'AMIA.CA', 'AMII.CA', 'AMOC.CA', 'AMPI.CA', 'APSW.CA', 'ARAB.CA', 'ARCC.CA',
    'AREH.CA', 'ASCM.CA', 'ASPI.CA', 'ATLC.CA', 'ATQA.CA', 'AXPH.CA', 'BIDI.CA',
    'BIGP.CA', 'BINV.CA', 'BIOC.CA', 'BONY.CA', 'BTFH.CA', 'CAED.CA', 'CANA.CA',
    'CCAP.CA', 'CCRS.CA', 'CEFM.CA', 'CERA.CA', 'CFGH.CA', 'CICH.CA', 'CIEB.CA',
    'CIRA.CA', 'CLHO.CA', 'CNFN.CA', 'COMI.CA', 'COPR.CA', 'COSG.CA', 'CPCI.CA',
    'CPME.CA', 'CRST.CA', 'CSAG.CA', 'DAPH.CA', 'DCRC.CA', 'DEIN.CA', 'DGTZ.CA',
    'DOMT.CA', 'DSCW.CA', 'DTPP.CA', 'EALR.CA', 'EASB.CA', 'EAST.CA', 'EBSC.CA',
    'ECAP.CA', 'EDFM.CA', 'EEII.CA', 'EFIC.CA', 'EFID.CA', 'EFIH.CA', 'EGAL.CA',
    'EGAS.CA', 'EGBE.CA', 'EGCH.CA', 'EGREF.CA', 'EGSA.CA', 'EGTS.CA', 'EHDR.CA',
    'ELAB.CA', 'ELEC.CA', 'ELKA.CA', 'ELNA.CA', 'ELSH.CA', 'ELWA.CA', 'EMFD.CA',
    'ENGC.CA', 'EOSB.CA', 'EPCO.CA', 'EPPK.CA', 'ETEL.CA', 'ETRS.CA', 'EXPA.CA',
    'FAIT.CA', 'FAITA.CA', 'FCMD.CA', 'FIRE.CA', 'FNAR.CA', 'FTNS.CA', 'FWRY.CA',
    'GBCO.CA', 'GDWA.CA', 'GGCC.CA', 'GGRN.CA', 'GIHD.CA', 'GMCI.CA', 'GOUR.CA',
    'GPIM.CA', 'GRCA.CA', 'GSSC.CA', 'GTEX.CA', 'GTHE.CA', 'GTWL.CA', 'HBCO.CA',
    'HDBK.CA', 'HELI.CA', 'HRHO.CA', 'IBCT.CA', 'ICFC.CA', 'ICID.CA', 'IDRE.CA',
    'IEEC.CA', 'IFAP.CA', 'INEG.CA', 'INFI.CA', 'IRON.CA', 'ISMA.CA', 'ISMQ.CA',
    'ISPH.CA', 'JUFO.CA', 'KABO.CA', 'KORA.CA', 'KRDI.CA', 'KWIN.CA', 'KZPC.CA',
    'LCSW.CA', 'LKGP.CA', 'LUTS.CA', 'MAAL.CA', 'MASR.CA', 'MBEG.CA', 'MBSC.CA',
    'MCQE.CA', 'MCRO.CA', 'MENA.CA', 'MEPA.CA', 'MFPC.CA', 'MFSC.CA', 'MHOT.CA',
    'MICH.CA', 'MILS.CA', 'MIPH.CA', 'MOED.CA', 'MOIL.CA', 'MOIN.CA', 'MOSC.CA',
    'MPCI.CA', 'MPCO.CA', 'MPRC.CA', 'MTIE.CA', 'NAHO.CA', 'NARE.CA', 'NCCW.CA',
    'NCGC.CA', 'NEDA.CA', 'NHPS.CA', 'NINH.CA', 'NIPH.CA', 'OBRI.CA', 'OCAP.CA',
    'OCDI.CA', 'OCPH.CA', 'ODIN.CA', 'OFH.CA', 'OIH.CA', 'OLFI.CA', 'ORAS.CA',
    'ORHD.CA', 'ORWE.CA', 'PHAR.CA', 'PHDC.CA', 'PHGC.CA', 'PHTV.CA', 'POUL.CA',
    'PRCL.CA', 'PRDC.CA', 'PRMH.CA', 'QNBE.CA', 'RACC.CA', 'RAKT.CA', 'RAYA.CA',
    'RKAZ.CA', 'RMDA.CA', 'RMTV.CA', 'ROTO.CA', 'RREI.CA', 'RTVC.CA', 'RUBX.CA',
    'SAUD.CA', 'SCEM.CA', 'SCFM.CA', 'SCTS.CA', 'SDTI.CA', 'SEIG.CA', 'SIEG.CA',
    'SIPC.CA', 'SKPC.CA', 'SMFR.CA', 'SNFC.CA', 'SPIN.CA', 'SPMD.CA', 'SUCE.CA',
    'SUGR.CA', 'SVCE.CA', 'SWDY.CA', 'TALM.CA', 'TANM.CA', 'TAQA.CA', 'TMGH.CA',
    'TORA.CA', 'TWSA.CA', 'TYCN.CA', 'UBEE.CA', 'UEFM.CA', 'UEGC.CA', 'UNIP.CA',
    'UNIT.CA', 'UPMS.CA', 'UTOP.CA', 'VALU.CA', 'VERT.CA', 'VLMR.CA', 'VLMRA.CA',
    'WCDF.CA', 'WKOL.CA', 'ZEOT.CA', 'ZMID.CA',
]


# ====================================================================
# 4. BACKTEST ENGINE & PERFORMANCE CALCULATIONS
# ====================================================================

class Backtester:

    def __init__(self, tickers: List[str], years: int = 10):
        self.tickers = tickers
        self.years = years
        self.detector = EveEveDetector()
        self.trades = []

    def run(self):
        end_date = datetime.now()
        start_date = end_date - timedelta(days=365 * self.years)

        print(f"Fetching TradingView real-time prices for EGX stocks...")
        tv_live_dict = fetch_tradingview_live_data(self.tickers)

        print(
            f"Starting Backtest & Pattern Scan for {self.years} YEARS "
            f"({start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')})...\n"
        )

        for idx, ticker in enumerate(self.tickers, 1):
            print(f'[{idx}/{len(self.tickers)}] Processing Symbol: {ticker} ...')
            try:
                df = get_combined_stock_data(ticker, start_date, end_date, tv_live_dict)

                if df.empty or len(df) < 50:
                    continue

                patterns = self.detector.detect(df)

                for p in patterns:
                    self._simulate_trade(df, ticker, p)

            except Exception as e:
                print(f'     -> Error processing {ticker}: {e}')

        self._print_performance_summary()

    def _simulate_trade(self, df: pd.DataFrame, ticker: str, p: EveEvePattern):
        entry_idx = p.breakout_index
        if entry_idx is None or entry_idx >= len(df):
            return

        entry_date = df.index[entry_idx]
        entry_price = p.entry_price
        target_price = p.target_price
        stop_loss_price = p.stop_loss

        exit_date = None
        status = 'Open'
        current_price = df['Close'].iloc[-1]

        for i in range(entry_idx + 1, len(df)):
            current_high = df['High'].iloc[i]
            current_low = df['Low'].iloc[i]
            c_date = df.index[i]

            if current_low <= stop_loss_price:
                exit_date = c_date
                status = 'Loss'
                break
            elif current_high >= target_price:
                exit_date = c_date
                status = 'Win'
                break

        exit_date_str = exit_date.strftime('%Y-%m-%d') if exit_date else ""

        self.trades.append({
            'Stock Name': ticker,
            'Entry Date': entry_date.strftime('%Y-%m-%d'),
            'Entry Price': round(float(entry_price), 2),
            'Target': round(float(target_price), 2),
            'Stop Loss': round(float(stop_loss_price), 2),
            'Exit Date': exit_date_str,
            'Current Price': round(float(current_price), 2),
            'Status': status,
        })

    def _print_performance_summary(self):
        if not self.trades:
            print("\n--------------------------------------------------------")
            print("لم يتم العثور على أي صفقات مطابقة لشروط النموذج خلال هذه الفترة.")
            print("--------------------------------------------------------")
            return

        df_trades = pd.DataFrame(self.trades)
        df_trades.drop_duplicates(subset=['Stock Name', 'Entry Date'], inplace=True)

        total_trades = len(df_trades)
        wins = df_trades[df_trades['Status'] == 'Win']
        losses = df_trades[df_trades['Status'] == 'Loss']
        opens = df_trades[df_trades['Status'] == 'Open']

        total_wins = len(wins)
        total_losses = len(losses)
        total_opens = len(opens)

        closed_trades = total_wins + total_losses

        win_rate = (total_wins / closed_trades * 100) if closed_trades > 0 else 0.0
        loss_rate = (total_losses / closed_trades * 100) if closed_trades > 0 else 0.0

        avg_win_pct = (((wins['Target'] - wins['Entry Price']) / wins['Entry Price']) * 100).mean() if not wins.empty else 0.0
        avg_loss_pct = (((losses['Entry Price'] - losses['Stop Loss']) / losses['Entry Price']) * 100).mean() if not losses.empty else 0.0

        print("\n========================================================")
        print(f"       EVE & EVE 10-YEAR BACKTEST PERFORMANCE           ")
        print("========================================================")
        print(f" إجمالي عدد الصفقات المكتشفة     : {total_trades}")
        print(f" الصفقات المغلقة (Win + Loss)  : {closed_trades}")
        print(f" الصفقات المفتوحة (حتى الآن)  : {total_opens}")
        print("--------------------------------------------------------")
        print(f" عدد الصفقات الناجحة (Win)     : {total_wins}")
        print(f" عدد الصفقات الفاشلة (Loss)    : {total_losses}")
        print("--------------------------------------------------------")
        print(f" نسبة النجاح (Win Rate)         : {win_rate:.2f}%")
        print(f" نسبة الفشل (Loss Rate)         : {loss_rate:.2f}%")
        print("--------------------------------------------------------")
        print(f" متوسط نسبة الربح في الصفقات الناجحة : +{avg_win_pct:.2f}%")
        print(f" متوسط نسبة الخسارة في الصفقات الفاشلة: -{avg_loss_pct:.2f}%")
        
        if avg_loss_pct > 0:
            rr_ratio = avg_win_pct / avg_loss_pct
            print(f" نسبة العائد إلى المخاطرة (Reward/Risk): {rr_ratio:.2f}")

        print("========================================================\n")

    def export_to_excel(self, filename: str = 'EGX_Eve_Eve_10Years_Results.xlsx'):
        if not self.trades:
            return

        trades_df = pd.DataFrame(self.trades)
        trades_df.drop_duplicates(subset=['Stock Name', 'Entry Date'], inplace=True)
        
        trades_df['Temp_Entry'] = pd.to_datetime(trades_df['Entry Date'])
        trades_df.sort_values(by='Temp_Entry', ascending=True, inplace=True)
        trades_df.drop(columns=['Temp_Entry'], inplace=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Backtest Results'
        ws.views.sheetView[0].rightToLeft = False

        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
        normal_font = Font(name='Calibri', size=11, bold=False, color='000000')

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9'),
        )

        align_center = Alignment(horizontal='center', vertical='center')

        columns = [
            'Stock Name',
            'Entry Date',
            'Entry Price',
            'Target',
            'Stop Loss',
            'Exit Date',
            'Current Price',
            'Status'
        ]
        ws.append(columns)

        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

        for r_idx, row in trades_df.iterrows():
            row_data = [row[col] for col in columns]
            ws.append(row_data)
            row_num = ws.max_row

            for c_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_num, column=c_idx)
                cell.border = thin_border
                cell.alignment = align_center
                cell.font = normal_font

                col_name = columns[c_idx - 1]

                if col_name in ['Entry Price', 'Target', 'Stop Loss', 'Current Price']:
                    cell.number_format = '#,##0.00'

        col_widths = {
            'A': 18, 'B': 16, 'C': 15, 'D': 15,
            'E': 15, 'F': 16, 'G': 16, 'H': 14
        }
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

        try:
            wb.save(filename)
            print(f'Excel File Saved Successfully: {filename}')
        except PermissionError:
            alt_filename = f"EGX_Eve_Eve_10Y_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            wb.save(alt_filename)
            print(f"[Permission Error] Main file was locked. Saved as: {alt_filename}")

        if COLAB_ENV:
            files.download(filename)


# ====================================================================
# 5. RUN MAIN PROGRAM
# ====================================================================

if __name__ == '__main__':
    # تشغيل الفحص لنموذج Eve & Eve على مدار 10 سنوات
    backtester = Backtester(tickers=EGYPTIAN_STOCKS, years=10)
    backtester.run()
    backtester.export_to_excel()
